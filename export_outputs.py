"""Create combined PDF, Word summary and Excel offer comparison from an app backup."""
from pathlib import Path
import json, sys
from pypdf import PdfReader, PdfWriter
from docx import Document
from openpyxl import Workbook
from template_renderer import from_app_state

def export_all(backup, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    tmpdir = output_dir / '.rendered'
    tmpdir.mkdir(parents=True, exist_ok=True)
    from_app_state(backup, tmpdir)
    writer = PdfWriter()
    for pdf in sorted(tmpdir.glob('*.pdf')):
        for page in PdfReader(str(pdf)).pages:
            writer.add_page(page)
    writer.write(str(output_dir / 'dogrudan-temin-dosyasinin-tamami.pdf'))
    for pdf in tmpdir.glob('*.pdf'): pdf.unlink()
    tmpdir.rmdir()

    p = (backup.get('procurements') or [{}])[-1]
    institution = backup.get('institution', {})
    doc = Document()
    doc.add_heading('Doğrudan Temin Dosyası', 0)
    doc.add_paragraph(f"Kurum: {institution.get('name','')}")
    doc.add_paragraph(f"İşin adı: {p.get('workName','')}")
    doc.add_paragraph(f"Tarih: {p.get('date','')}    Hukuki dayanak: {p.get('law','')}")
    doc.add_heading('Kalemler ve teklifler', level=1)
    companies = backup.get('companies', [])
    items = p.get('items', [])
    table = doc.add_table(rows=1, cols=4 + len(companies))
    headers = ['Sıra', 'Açıklama', 'Miktar', 'Birim'] + [c.get('name','') for c in companies]
    for cell, value in zip(table.rows[0].cells, headers): cell.text = str(value)
    for n, item in enumerate(items, 1):
        cells = table.add_row().cells
        for cell, value in zip(cells, [n, item.get('description',''), item.get('quantity',''), item.get('unit','')] + [x or '' for x in item.get('prices', [])]): cell.text = str(value)
    doc.save(str(output_dir / 'dogrudan-temin-dosyasi.docx'))

    wb = Workbook(); ws = wb.active; ws.title = 'Teklif Karşılaştırma'
    ws.append(['Sıra', 'Kalem', 'Miktar', 'Birim'] + [c.get('name','') for c in companies])
    for n, item in enumerate(items, 1): ws.append([n, item.get('description',''), item.get('quantity',''), item.get('unit','')] + [x or '' for x in item.get('prices', [])])
    ws.append([]); ws.append(['Firma toplamları'])
    for ci, company in enumerate(companies, 5):
        ws.cell(ws.max_row, ci).value = f'=SUMPRODUCT(C2:C{ws.max_row-2},{chr(64+ci)}2:{chr(64+ci)}{ws.max_row-2})'
    wb.save(str(output_dir / 'teklif-karsilastirma.xlsx'))

if __name__ == '__main__':
    if len(sys.argv) != 3: raise SystemExit('Kullanım: python export_outputs.py yedek.json çıktı-klasörü')
    export_all(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')), Path(sys.argv[2]))
